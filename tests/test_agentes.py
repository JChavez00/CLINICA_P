"""Proveedores simulados; SQLite aislado y sin datos reales de pacientes."""
import io
import json
import os
import unittest
from datetime import date, datetime, timezone, timedelta
from uuid import uuid4
from unittest.mock import patch

os.environ["DATABASE_URL"]="sqlite:///:memory:"
os.environ["SECRET_KEY"]="isolated-test-session"
os.environ["AGENTES_POR_MINUTO"]="100"

from flask import Flask
from sqlalchemy.schema import CreateTable
from sqlalchemy.dialects import postgresql
from openpyxl import load_workbook

from models import (db,Paciente,Atencion,AgenteEjecucion,InformeIA,TareaIA,TareaIAEjecucion)
from llm.rutas import bp_asistente
from llm.esquemas import ErrorAgente,Reporte
from llm.privacidad import clasificar_chat
from llm.tareas import ejecutar_programadas,confirmar


def modelo_simulado(proveedor,instrucciones,datos,esquema):
    if "operacion_sugerida" in datos:
        return {"herramienta":datos["operacion_sugerida"],"argumentos":datos["parametros_validados"]}
    return {"hechos":list(datos["hechos_disponibles"]),"acciones":[]}


class AgentesTests(unittest.TestCase):
    def setUp(self):
        from pathlib import Path
        raiz=Path(__file__).resolve().parent.parent
        self.app=Flask("test_agentes",template_folder=str(raiz/"templates"))
        self.app.config.update(SQLALCHEMY_DATABASE_URI="sqlite:///:memory:",SECRET_KEY="test",TESTING=True)
        db.init_app(self.app); self.app.register_blueprint(bp_asistente)
        # Endpoints usados por la plantilla de navegación, sin importar app productiva.
        for endpoint in ["dashboard","login","nueva_atencion","pacientes_hoy","reportes","logout","demanda.pagina"]:
            self.app.add_url_rule('/stub/'+endpoint,endpoint,lambda:"stub")
        self.ctx=self.app.app_context();self.ctx.push();db.create_all()
        p=Paciente(dni="12345678",nombres="ALICIA",apellidos="PÉREZ",telefono="999888777",nombre_via="DOMICILIO PRIVADO",fecha_nacimiento="2000-01-01")
        p2=Paciente(dni="87654321",nombres="ALICIA",apellidos="PEREZ",fecha_nacimiento="2000-01-01")
        p3=Paciente(dni="INVALIDO",nombres="",apellidos="")
        db.session.add_all([p,p2,p3]);db.session.flush()
        self.pid=p.id
        db.session.add_all([Atencion(paciente_id=p.id,fecha=date(2026,9,1),motivo="CLINICO SECRETO",diagnostico="DIAGNOSTICO PRIVADO",estado="Finalizado"),
                            Atencion(paciente_id=p2.id,fecha=date(2026,9,2),motivo="CONTROL",estado="Finalizado"),
                            Atencion(paciente_id=p3.id,fecha=date(2026,9,3),motivo="",estado="OTRO")]);db.session.commit()
        self.c=self.app.test_client()
        with self.c.session_transaction() as s:s['autenticado']=True;s['ia_csrf']='csrf-test'
        self.headers={"X-CSRF-Token":"csrf-test"}
        from llm.clientes import solicitar
        self.solicitar_real=solicitar
        self.mock=patch('llm.clientes.solicitar',side_effect=modelo_simulado)
        self.llm=self.mock.start()
    def tearDown(self):
        self.mock.stop();db.session.remove();self.ctx.pop()
    def post(self,url,datos): return self.c.post(url,json=datos,headers=self.headers)

    def test_seleccion_agente(self):
        self.assertEqual(clasificar_chat('Busca DNI 12345678')['agente'],'consultor')
        self.assertEqual(clasificar_chat('Cuantas atenciones hubo en septiembre')['agente'],'analista')
        self.assertIsNone(clasificar_chat('ALICIA tiene fiebre en domicilio privado'))

    def test_dni_y_contexto(self):
        r=self.post('/api/asistente-ia',{'pregunta':'Busca DNI 12345678'})
        self.assertEqual(r.status_code,200);self.assertIn('ALICIA',r.json['respuesta'])
        self.assertEqual(r.json['proveedor'],'gemini')
        r=self.post('/api/asistente-ia',{'pregunta':'Su historial'})
        self.assertIn('DIAGNOSTICO PRIVADO',r.json['respuesta'])

    def test_nombres_multiples_seleccion(self):
        r=self.post('/api/asistente-ia',{'pregunta':'Busca al paciente Alicia Perez'})
        self.assertEqual(len(r.json['opciones']),2)
        elegido=self.post('/api/asistente-ia',{'seleccion':self.pid})
        self.assertEqual(elegido.status_code,200);self.assertIn('DIAGNOSTICO PRIVADO',elegido.json['respuesta'])
        self.assertEqual(self.post('/api/asistente-ia',{'seleccion':99999}).status_code,403)

    def test_paciente_inexistente(self):
        r=self.post('/api/asistente-ia',{'pregunta':'Busca DNI 11111111'})
        self.assertIn('No encontré',r.json['respuesta'])

    def test_contexto_no_se_reutiliza_para_nombre_ambiguo(self):
        self.post('/api/asistente-ia',{'pregunta':'Busca DNI 12345678'})
        r=self.post('/api/asistente-ia',{'pregunta':'El diagnostico de otra persona'})
        self.assertNotIn('DIAGNOSTICO PRIVADO',r.json['respuesta'])

    def test_privacidad_frontera_y_auditoria(self):
        self.post('/api/asistente-ia',{'pregunta':'Busca DNI 12345678 con diagnostico DIAGNOSTICO PRIVADO'})
        self.post('/api/asistente-ia',{'pregunta':'Cuantas atenciones hubo en septiembre de ALICIA telefono 999888777'})
        enviados=json.dumps(self.llm.call_args_list,default=str)
        for secreto in ['12345678','999888777','ALICIA','DIAGNOSTICO PRIVADO','DOMICILIO PRIVADO','CLINICO SECRETO']:
            self.assertNotIn(secreto,enviados)
        logs=json.dumps(self.c.get('/api/agentes/ejecuciones').json)
        self.assertNotIn('12345678',logs);self.assertNotIn('ALICIA',logs)

    def test_estadisticas_agregadas(self):
        r=self.post('/api/asistente-ia',{'pregunta':'Cuantas atenciones hubo en septiembre'})
        self.assertEqual(r.json['agente'],'analista_inteligente')
        self.assertEqual(r.json['datos']['total_atenciones'],3)
        self.assertTrue(all(c.args[0]=='groq' for c in self.llm.call_args_list))

    def test_comparacion(self):
        r=self.post('/api/asistente-ia',{'pregunta':'Compara agosto y septiembre'})
        self.assertEqual(r.json['datos']['diferencia'],3)

    def test_campos_no_inventados(self):
        r=self.post('/api/asistente-ia',{'pregunta':'Cuantos servicios y costos hubo en septiembre'})
        self.assertFalse(r.json['datos']['disponible'])

    def test_chat_no_invoca_otros_modulos(self):
        from llm.coordinador import coordinar
        with self.assertRaises(ErrorAgente): coordinar('chat','tareas',{})

    def informe(self):
        r=self.post('/api/agentes/reportes',{'configuracion':{'desde':'2026-09-01','hasta':'2026-09-30','tipo':'mensual','secciones':['atenciones','servicios','costos']}})
        self.assertEqual(r.status_code,201,r.json)
        return r.json

    def test_reporte_y_cifras(self):
        d=self.informe();self.assertEqual(d['datos']['estadisticas']['total_atenciones'],3)
        self.assertIn('no_disponible',d['datos']);self.assertEqual(InformeIA.query.count(),1)

    def test_secciones_no_seleccionadas(self):
        r=self.post('/api/agentes/reportes',{'configuracion':{'desde':'2026-09-01','hasta':'2026-09-30','secciones':['costos']}})
        self.assertEqual(r.status_code,201)
        self.assertIsNone(r.json['datos']['estadisticas']['total_atenciones'])
        self.assertIn('no_disponible',r.json['datos'])

    def test_error_cliente_no_expone_mensaje_del_proveedor(self):
        with patch.dict(os.environ,{'GROQ_API_KEY':'clave_simulada'}), patch('llm.clientes._cliente') as cliente:
            cliente.return_value.chat.completions.create.side_effect=RuntimeError('SECRETO DEL PROVEEDOR')
            with self.assertRaises(ErrorAgente) as error: self.solicitar_real('groq','Contrato',{}, {})
            self.assertNotIn('SECRETO',str(error.exception))

    def test_exportacion_pdf(self):
        d=self.informe();r=self.c.get(f"/api/agentes/reportes/{d['id']}/pdf")
        self.assertEqual(r.status_code,200);self.assertTrue(r.data.startswith(b'%PDF'))
        self.assertGreater(len(r.data),2000)

    def test_exportacion_excel(self):
        d=self.informe();r=self.c.get(f"/api/agentes/reportes/{d['id']}/xlsx")
        wb=load_workbook(io.BytesIO(r.data));self.assertEqual(wb['Resumen']['B5'].value,3)
        self.assertEqual(wb['Atenciones'].max_row,4);self.assertTrue(wb['Mensual']._charts)

    def test_colaboracion_opcional(self):
        r=self.post('/api/agentes/reportes',{'configuracion':{'desde':'2026-09-01','hasta':'2026-09-30'},'revision_groq':True})
        self.assertEqual(r.status_code,201);self.assertIn('revision_adicional',r.json['datos'])
        self.assertEqual(set(c.args[0] for c in self.llm.call_args_list),{'gemini','groq'})

    def test_tarea_exige_confirmacion_y_csrf(self):
        self.assertEqual(self.c.post('/api/agentes/tareas/preview',json={'solicitud':'Reporte diario'}).status_code,403)
        r=self.post('/api/agentes/tareas/preview',{'solicitud':'Genera un reporte todos los lunes a las 08:00'})
        self.assertEqual(r.status_code,200);self.assertEqual(TareaIA.query.count(),0)
        token=r.json['confirmacion_token']
        self.assertEqual(self.post('/api/agentes/tareas/confirmar',{'token':token}).status_code,400)
        confirmado=self.post('/api/agentes/tareas/confirmar',{'token':token,'confirmar':True})
        self.assertEqual(confirmado.status_code,201)
        self.post('/api/agentes/tareas/confirmar',{'token':token,'confirmar':True})
        self.assertEqual(TareaIA.query.count(),1)

    def test_token_tarea_no_se_puede_falsificar(self):
        self.assertEqual(self.post('/api/agentes/tareas/confirmar',{'token':'alterado','confirmar':True}).status_code,400)

    def test_destinatarios_no_autorizados(self):
        r=self.post('/api/agentes/tareas/preview',{'configuracion':{'tipo':'reporte','frecuencia':'diaria','canal':'email','destinatario':'tercero@example.com'}})
        self.assertEqual(r.status_code,400)

    def test_tarea_desconocida_no_se_programa(self):
        r=self.post('/api/agentes/tareas/preview',{'solicitud':'Elimina todos los pacientes'})
        self.assertEqual(r.status_code,400);self.assertEqual(TareaIA.query.count(),0)

    def test_alertas_sin_emisiones_no_inventan_demanda(self):
        from llm.tareas import ejecutar_una
        tarea=self.crear_tarea('alerta_alta')
        with patch('llm.herramientas.consultar_pronosticos_ml',return_value={'disponible':False}):
            r=ejecutar_una(tarea,date(2026,10,22))
        self.assertEqual(r['emisiones_evaluadas'],0);self.assertEqual(r['alertas'],[]);self.assertFalse(r['disponible'])

    def crear_tarea(self,tipo='estadisticas'):
        tarea=confirmar({'tipo':tipo,'frecuencia':'diaria','hora':0},str(uuid4()))
        # El reloj de estas pruebas es 22/10/2026: crear antes de la ocurrencia,
        # ya que las tareas nuevas ahora programan la primera hora futura.
        from models import TareaPerfil
        tarea.creada_en=datetime(2026,10,21,12)
        db.session.get(TareaPerfil,tarea.id).proxima=datetime(2026,10,22,5)
        db.session.commit()
        return tarea

    def test_cron_idempotente(self):
        t=self.crear_tarea();ahora=datetime(2026,10,22,10,tzinfo=timezone(timedelta(hours=-5)))
        self.assertEqual(len(ejecutar_programadas(ahora)),1)
        self.assertEqual(ejecutar_programadas(ahora),[])
        self.assertEqual(TareaIAEjecucion.query.count(),1)

    def test_reintentos_limitados(self):
        self.crear_tarea();ahora=datetime(2026,10,22,10,tzinfo=timezone(timedelta(hours=-5)))
        with patch('llm.tareas.ejecutar_una',side_effect=ErrorAgente('proveedor_no_disponible')):
            for _ in range(4): ejecutar_programadas(ahora)
        e=TareaIAEjecucion.query.one();self.assertEqual(e.intentos,3);self.assertEqual(e.estado,'error')

    def test_ejecutar_tarea_confirmada_desactivar(self):
        t=self.crear_tarea()
        self.assertEqual(self.post(f'/api/agentes/tareas/{t.id}/ejecutar',{'ejecutar':True}).status_code,200)
        self.assertEqual(self.post(f'/api/agentes/tareas/{t.id}',{'activa':False}).status_code,200)
        self.assertEqual(self.post(f'/api/agentes/tareas/{t.id}/ejecutar',{'ejecutar':True}).status_code,400)

    def test_calidad_detecta_incompletos_y_sospechas(self):
        r=self.post('/api/agentes/calidad',{'desde':'2026-09-01','hasta':'2026-09-30'})
        self.assertEqual(r.status_code,200,r.json)
        a=r.json['datos']['agregado'];self.assertEqual(a['posibles_duplicados'],2)
        self.assertIn('campos_obligatorios',a['categorias']);self.assertIn('atencion_incompleta',a['categorias'])
        self.assertEqual(Paciente.query.count(),3)
        enviados=json.dumps(self.llm.call_args_list,default=str)
        self.assertNotIn('ALICIA',enviados);self.assertNotIn('incidencias_locales',enviados)

    def test_fallo_gemini_consulta_local(self):
        self.llm.side_effect=ErrorAgente('modelo_no_disponible',503)
        r=self.post('/api/asistente-ia',{'pregunta':'Busca DNI 12345678'})
        self.assertEqual(r.status_code,200);self.assertEqual(r.json['estado'],'degradado');self.assertIn('ALICIA',r.json['respuesta'])

    def test_fallo_groq_estadisticas_locales(self):
        self.llm.side_effect=ErrorAgente('limite_proveedor',503)
        r=self.post('/api/asistente-ia',{'pregunta':'Cuantas atenciones hubo en septiembre'})
        self.assertEqual(r.status_code,200);self.assertEqual(r.json['datos']['total_atenciones'],3)
        self.assertEqual(r.json['estado'],'degradado')

    def test_plan_malicioso_no_ejecuta_herramientas(self):
        self.llm.return_value={'herramienta':'ejecutar_sql','argumentos':{'sql':'DELETE'}};self.llm.side_effect=None
        r=self.post('/api/asistente-ia',{'pregunta':'Cuantas atenciones hubo en septiembre'})
        self.assertEqual(r.json['estado'],'degradado');self.assertEqual(Atencion.query.count(),3)

    def test_plan_no_amplia_periodo(self):
        self.llm.side_effect=lambda *a:{'herramienta':'estadisticas_periodo','argumentos':{'desde':'2000-01-01','hasta':'2026-09-30'}}
        r=self.post('/api/asistente-ia',{'pregunta':'Cuantas atenciones hubo en septiembre'})
        self.assertEqual(r.json['datos']['desde'],'2026-09-01');self.assertEqual(r.json['estado'],'degradado')

    def test_conclusiones_no_inventan_cifras(self):
        def mock(p,i,d,e):
            if 'operacion_sugerida' in d:return modelo_simulado(p,i,d,e)
            return {'hechos':['cifra_inventada_99999'],'acciones':[]}
        self.llm.side_effect=mock
        r=self.post('/api/asistente-ia',{'pregunta':'Cuantas atenciones hubo en septiembre'})
        self.assertNotIn('99999',r.json['respuesta']);self.assertEqual(r.json['estado'],'degradado')

    def test_permisos_y_anonimo(self):
        self.app.config['AGENTES_PERMISOS']=['estadisticas']
        self.assertEqual(self.post('/api/asistente-ia',{'pregunta':'Busca DNI 12345678'}).status_code,403)
        self.assertEqual(self.c.get('/api/agentes/ejecuciones').status_code,403)
        c=self.app.test_client();self.assertEqual(c.get('/api/agentes/tareas').status_code,401)

    def test_limites_y_fechas(self):
        self.assertEqual(self.post('/api/asistente-ia',{'pregunta':'x'*501}).status_code,400)
        self.assertEqual(self.post('/api/agentes/calidad',{'desde':'2026-02-30','hasta':'2026-09-30'}).status_code,400)
        self.assertEqual(self.post('/api/agentes/calidad',{'desde':'2020-01-01','hasta':'2026-09-30'}).status_code,400)

    def test_postgres_ddl(self):
        for m in [AgenteEjecucion,InformeIA,TareaIA,TareaIAEjecucion]:
            self.assertIn(m.__tablename__,str(CreateTable(m.__table__).compile(dialect=postgresql.dialect())))

    def test_paginas_modulos(self):
        for ruta in ['/asistente-ia','/reportes-inteligentes','/automatizacion-tareas','/calidad-datos','/actividad-agentes']:
            self.assertEqual(self.c.get(ruta).status_code,200,ruta)


if __name__=='__main__':unittest.main()
