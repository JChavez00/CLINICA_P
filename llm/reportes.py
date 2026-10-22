"""Documentos con cifras y gráficos deterministas, nunca texto libre del proveedor."""
import io
from datetime import date, timezone, timedelta
from uuid import uuid4
from xml.sax.saxutils import escape

import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment

from models import db, InformeIA, ahora_utc


def guardar(resultado,identificador=None):
    identificador=identificador or str(uuid4())
    existente=db.session.get(InformeIA,identificador)
    if existente: return existente
    datos=resultado["datos"]
    informe=InformeIA(id=identificador,desde=date.fromisoformat(datos["periodo"]["desde"]),
                     hasta=date.fromisoformat(datos["periodo"]["hasta"]),
                     contenido={"datos":datos,"estado":resultado["estado"],"agente":resultado["agente"],
                                "advertencia":resultado.get("advertencia")})
    db.session.add(informe); db.session.commit()
    from .almacenamiento import conservar
    conservar(informe)
    return informe


def publico(informe):
    return {"id":informe.id,"generado_en":informe.generado_en.isoformat()+"Z",
            "desde":informe.desde.isoformat(),"hasta":informe.hasta.isoformat(),**informe.contenido}


def hora_peru(valor):
    return valor.replace(tzinfo=timezone.utc).astimezone(timezone(timedelta(hours=-5))).replace(tzinfo=None)


def excel(informe):
    wb=openpyxl.Workbook(); ws=wb.active; ws.title="Resumen"
    datos=informe.contenido["datos"]; e=datos["estadisticas"]
    ws.append(["CLINICA_P", "Reporte administrativo"])
    ws.append(["Desde",informe.desde]); ws.append(["Hasta",informe.hasta])
    ws.append(["Generado (Perú)",hora_peru(informe.generado_en)])
    ws.append(["Atenciones",e["total_atenciones"]]); ws.append(["Pacientes distintos",e["pacientes_unicos"]])
    ws.append(["Estado IA",informe.contenido["estado"]])
    for c in datos.get("conclusiones",[]): ws.append(["Conclusión",c])
    if "no_disponible" in datos: ws.append(["Campos no disponibles",datos["no_disponible"]["mensaje"]])
    diario=wb.create_sheet("Atenciones")
    diario.append(["Fecha observada","Atenciones"])
    for d in e["por_dia"]: diario.append([date.fromisoformat(d["fecha"]),d["cantidad"]])
    meses=wb.create_sheet("Mensual"); meses.append(["Mes","Atenciones"])
    for m in e["por_mes"]: meses.append([m["mes"],m["cantidad"]])
    if e["por_mes"]:
        chart=BarChart(); chart.title="Atenciones registradas"; chart.y_axis.title="Atenciones"
        chart.add_data(Reference(meses,min_col=2,min_row=1,max_row=meses.max_row),titles_from_data=True)
        chart.set_categories(Reference(meses,min_col=1,min_row=2,max_row=meses.max_row)); meses.add_chart(chart,"D2")
    if "ml" in datos:
        ml=wb.create_sheet("Pronosticos ML"); ml.append(["Fecha","Tipo","Demanda","Estimado","Real"])
        for d in datos["ml"].get("resultados",[]): ml.append([date.fromisoformat(d["fecha"]),d["tipo"],d["nivel"],d["cantidad"],d["real"]])
        ml.append(["Advertencia","Pronósticos experimentales; consultar referencias y métricas."])
        for objetivo,metricas in datos["ml"].get("modelos",{}).items():
            ml.append(["Objetivo",objetivo,"Modelo",metricas["ganador"]])
            ml.append(["Mejora frente a referencia", "Sí" if metricas["mejora_demostrada"] else "No demostrada"])
        for advertencia in datos["ml"].get("advertencias",[]): ml.append(["Limitación",advertencia])
    if "revision_adicional" in datos:
        for conclusion in datos["revision_adicional"]["conclusiones"]: ws.append(["Revisión adicional",conclusion])
    for hoja in wb:
        hoja.freeze_panes="A2"; hoja.auto_filter.ref=hoja.dimensions
        for c in hoja[1]: c.fill=PatternFill("solid",fgColor="1B4B43"); c.font=Font(color="FFFFFF",bold=True)
        for fila in hoja:
            for c in fila:
                c.alignment=Alignment(vertical="top",wrap_text=True)
                if isinstance(c.value,date): c.number_format="dd/mm/yyyy"
                if isinstance(c.value,str) and c.value.startswith(("=","+","-","@")): c.value="'"+c.value
        hoja.column_dimensions["A"].width=28; hoja.column_dimensions["B"].width=78 if hoja.title=="Resumen" else 30
        for col in ["C","D","E"]: hoja.column_dimensions[col].width=22
        for i in range(8,hoja.max_row+1): hoja.row_dimensions[i].height=45 if hoja.title=="Resumen" else 20
    salida=io.BytesIO(); wb.save(salida); salida.seek(0); return salida


def pdf(informe):
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether
    from reportlab.graphics.shapes import Drawing, String
    from reportlab.graphics.charts.barcharts import VerticalBarChart
    salida=io.BytesIO(); estilos=getSampleStyleSheet()
    datos=informe.contenido["datos"]; e=datos["estadisticas"]
    historia=[Paragraph("CLINICA_P · Reporte administrativo",estilos["Title"]),
              Paragraph(f"Período: {informe.desde.isoformat()} a {informe.hasta.isoformat()}",estilos["Normal"]),
              Paragraph(f"Generado: {hora_peru(informe.generado_en).strftime('%d/%m/%Y %H:%M')} (Perú)",estilos["Normal"]),Spacer(1,.5*cm)]
    tabla=Table([["Indicador","Cantidad"],["Atenciones",e["total_atenciones"] if e["total_atenciones"] is not None else "Sección no seleccionada"],["Pacientes distintos",e["pacientes_unicos"] if e["pacientes_unicos"] is not None else "—"]],colWidths=[11*cm,5*cm])
    tabla.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#1B4B43")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("BOTTOMPADDING",(0,0),(-1,-1),10),("GRID",(0,0),(-1,-1),.3,colors.lightgrey)]))
    historia.extend([tabla,Spacer(1,.5*cm)])
    if e["por_mes"]:
        dibujo=Drawing(450,210); barras=VerticalBarChart()
        barras.x=35; barras.y=35; barras.width=390; barras.height=145
        barras.data=[[d["cantidad"] for d in e["por_mes"]]]
        barras.categoryAxis.categoryNames=[d["mes"] for d in e["por_mes"]]
        barras.valueAxis.valueMin=0; barras.bars[0].fillColor=colors.HexColor("#1B4B43")
        barras.categoryAxis.labels.fontSize=8
        dibujo.add(barras); dibujo.add(String(35,195,"Atenciones observadas por mes",fontSize=11))
        historia.extend([dibujo,Spacer(1,.3*cm)])
    filas=[["Mes observado","Atenciones"]]+[[d["mes"],d["cantidad"]] for d in e["por_mes"]]
    historia.append(Table(filas,colWidths=[11*cm,5*cm],repeatRows=1))
    historia.extend([Spacer(1,.5*cm),Paragraph("Conclusiones",estilos["Heading2"])])
    for c in datos.get("conclusiones",[]): historia.extend([Paragraph(escape(c),estilos["BodyText"]),Spacer(1,.2*cm)])
    if "no_disponible" in datos: historia.append(Paragraph(escape(datos["no_disponible"]["mensaje"]),estilos["BodyText"]))
    if "ml" in datos:
        titulo_ml=Paragraph("Pronósticos ML existentes (experimentales)",estilos["Heading2"])
        filas=[["Fecha","Tipo","Demanda","Estimado","Real"]]
        for d in datos["ml"].get("resultados",[]): filas.append([d["fecha"],d["tipo"],d["nivel"] or "—",str(d["cantidad"] if d["cantidad"] is not None else "—"),str(d["real"] if d["real"] is not None else "—")])
        tabla_ml=Table(filas,repeatRows=1,colWidths=[90,90,80,80,80])
        tabla_ml.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#1B4B43")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("BOTTOMPADDING",(0,0),(-1,-1),7)]))
        historia.append(KeepTogether([titulo_ml,tabla_ml]))
        historia.append(Paragraph("Las predicciones se distinguen de las observaciones; no son certezas ni validación operativa.",estilos["BodyText"]))
        for objetivo,metricas in datos["ml"].get("modelos",{}).items():
            mejora="supera el criterio orientativo" if metricas["mejora_demostrada"] else "no demostró mejora suficiente frente a su referencia"
            historia.append(Paragraph(escape(f"{objetivo}: {metricas['ganador']} {mejora}. Sigue siendo experimental."),estilos["BodyText"]))
    if "revision_adicional" in datos:
        historia.append(Paragraph("Revisión adicional del analista",estilos["Heading2"]))
        for c in datos["revision_adicional"]["conclusiones"]: historia.append(Paragraph(escape(c),estilos["BodyText"]))
    if informe.contenido.get("advertencia"): historia.append(Paragraph(escape(informe.contenido["advertencia"]),estilos["BodyText"]))
    def pie(canvas,doc):
        canvas.setFont("Helvetica",8); canvas.drawString(2*cm,1.2*cm,"CLINICA_P · Datos agregados · Sin expedientes individuales")
        canvas.drawRightString(19*cm,1.2*cm,f"Página {doc.page}")
    SimpleDocTemplate(salida,rightMargin=2*cm,leftMargin=2*cm,topMargin=1.7*cm,bottomMargin=2*cm).build(historia,onFirstPage=pie,onLaterPages=pie)
    salida.seek(0); return salida
