import io
import os
from datetime import date, datetime, timedelta
from functools import wraps

from dotenv import load_dotenv
load_dotenv()  # carga variables desde .env si existe (solo afecta desarrollo local)

from flask import (
    Flask, render_template, request, redirect, url_for,
    session, flash, jsonify, send_file
)
from sqlalchemy import func

from config import Config
from models import db, Paciente, Atencion
from services.reniec import consultar_dni
from services.geocodificacion import buscar_direccion

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment


def create_app():
    app = Flask(__name__)
    app.config.from_object(Config)
    db.init_app(app)

    with app.app_context():
        db.create_all()

    # ---------------------------------------------------------------
    # Autenticación simple (un solo usuario: el doctor/administrador)
    # ---------------------------------------------------------------
    def login_requerido(vista):
        @wraps(vista)
        def envoltura(*args, **kwargs):
            if not session.get("autenticado"):
                return redirect(url_for("login"))
            return vista(*args, **kwargs)
        return envoltura

    @app.route("/login", methods=["GET", "POST"])
    def login():
        if session.get("autenticado"):
            return redirect(url_for("dashboard"))

        if request.method == "POST":
            usuario = request.form.get("usuario", "")
            clave = request.form.get("clave", "")
            if usuario == app.config["APP_USERNAME"] and clave == app.config["APP_PASSWORD"]:
                session.permanent = True
                session["autenticado"] = True
                session["usuario"] = usuario
                return redirect(url_for("dashboard"))
            flash("Usuario o clave incorrectos.", "error")

        return render_template("login.html")

    @app.route("/logout")
    def logout():
        session.clear()
        return redirect(url_for("login"))

    # ---------------------------------------------------------------
    # Dashboard
    # ---------------------------------------------------------------
    @app.route("/")
    @login_requerido
    def dashboard():
        hoy = date.today()
        atenciones_hoy = (
            Atencion.query.filter(Atencion.fecha == hoy)
            .order_by(Atencion.creado_en.desc())
            .all()
        )
        total_hoy = len(atenciones_hoy)
        en_proceso = sum(1 for a in atenciones_hoy if a.estado == "En proceso")

        hace_7_dias = hoy - timedelta(days=6)
        conteo_semana = (
            db.session.query(Atencion.fecha, func.count(Atencion.id))
            .filter(Atencion.fecha >= hace_7_dias, Atencion.fecha <= hoy)
            .group_by(Atencion.fecha)
            .all()
        )
        conteo_dict = {f.strftime("%Y-%m-%d"): c for f, c in conteo_semana}
        serie_semana = []
        for i in range(6, -1, -1):
            d = hoy - timedelta(days=i)
            serie_semana.append({
                "fecha": d.strftime("%d/%m"),
                "total": conteo_dict.get(d.strftime("%Y-%m-%d"), 0),
            })

        total_pacientes = Paciente.query.count()

        return render_template(
            "dashboard.html",
            total_hoy=total_hoy,
            en_proceso=en_proceso,
            total_pacientes=total_pacientes,
            atenciones_recientes=atenciones_hoy[:5],
            serie_semana=serie_semana,
            hoy=hoy,
        )

    # ---------------------------------------------------------------
    # Buscar DNI (AJAX) — revisa base local primero, luego API externa
    # ---------------------------------------------------------------
    @app.route("/api/buscar-dni/<dni>")
    # @login_requerido
    def buscar_dni(dni):
        dni = dni.strip()

        paciente = Paciente.query.filter_by(dni=dni).first()
        if paciente:
            return jsonify({
                "ok": True,
                "origen": "base_local",
                "nombres": paciente.nombres,
                "apellidos": paciente.apellidos,
                "fecha_nacimiento": paciente.fecha_nacimiento or "",
                "telefono": paciente.telefono or "",
                "sexo": paciente.sexo or "",
                "departamento": paciente.departamento or "",
                "provincia": paciente.provincia or "",
                "distrito": paciente.distrito or "",
                "tipo_via": paciente.tipo_via or "",
                "nombre_via": paciente.nombre_via or "",
                "latitud": paciente.latitud,
                "longitud": paciente.longitud,
                "es_nuevo": False,
            })

        resultado = consultar_dni(
            dni,
            app.config["DNI_API_URL"],
            app.config["DNI_API_TOKEN"],
        )

        if resultado.get("ok"):
            return jsonify({
                "ok": True,
                "origen": "api_reniec",
                "nombres": resultado["nombres"],
                "apellidos": resultado["apellidos"],
                "fecha_nacimiento": resultado.get("fecha_nacimiento", ""),
                "telefono": "",
                "sexo": "",
                "es_nuevo": True,
            })

        if resultado.get("error") == "no_configurado":
            return jsonify({
                "ok": False,
                "es_nuevo": True,
                "info": "La consulta automática a RENIEC no está configurada. Ingresa los datos manualmente.",
            })

        return jsonify({
            "ok": False,
            "es_nuevo": True,
            "info": resultado.get("error", "No se pudo consultar el DNI. Ingresa los datos manualmente."),
        })

    # ---------------------------------------------------------------
    # Buscar dirección (AJAX) — proxy hacia Nominatim/OpenStreetMap
    # ---------------------------------------------------------------
    @app.route("/api/buscar-direccion")
    # @login_requerido
    def buscar_direccion_ruta():
        consulta = request.args.get("q", "").strip()
        if not consulta:
            return jsonify({"ok": False, "error": "Escribe al menos parte de la dirección."})

        resultado = buscar_direccion(consulta)
        return jsonify(resultado)

    # ---------------------------------------------------------------
    # Nueva atención
    # ---------------------------------------------------------------
    @app.route("/nueva-atencion", methods=["GET", "POST"])
    @login_requerido
    def nueva_atencion():
        if request.method == "POST":
            dni = request.form.get("dni", "").strip()
            nombres = request.form.get("nombres", "").strip().upper()
            apellidos = request.form.get("apellidos", "").strip().upper()
            fecha_nacimiento = request.form.get("fecha_nacimiento", "").strip()
            telefono = request.form.get("telefono", "").strip()
            sexo = request.form.get("sexo", "").strip()

            # Domicilio
            departamento = request.form.get("departamento", "").strip().title()
            provincia = request.form.get("provincia", "").strip().title()
            distrito = request.form.get("distrito", "").strip().title()
            tipo_via = request.form.get("tipo_via", "").strip()
            nombre_via = request.form.get("nombre_via", "").strip()
            latitud = request.form.get("latitud", "").strip()
            longitud = request.form.get("longitud", "").strip()
            latitud = float(latitud) if latitud else None
            longitud = float(longitud) if longitud else None

            if not dni or not nombres or not apellidos:
                flash("DNI, nombres y apellidos son obligatorios.", "error")
                return redirect(url_for("nueva_atencion"))

            paciente = Paciente.query.filter_by(dni=dni).first()
            if paciente:
                paciente.nombres = nombres
                paciente.apellidos = apellidos
                if fecha_nacimiento:
                    paciente.fecha_nacimiento = fecha_nacimiento
                if telefono:
                    paciente.telefono = telefono
                if sexo:
                    paciente.sexo = sexo
                if departamento:
                    paciente.departamento = departamento
                if provincia:
                    paciente.provincia = provincia
                if distrito:
                    paciente.distrito = distrito
                if tipo_via:
                    paciente.tipo_via = tipo_via
                if nombre_via:
                    paciente.nombre_via = nombre_via
                if latitud is not None:
                    paciente.latitud = latitud
                if longitud is not None:
                    paciente.longitud = longitud
            else:
                paciente = Paciente(
                    dni=dni, nombres=nombres, apellidos=apellidos,
                    fecha_nacimiento=fecha_nacimiento, telefono=telefono, sexo=sexo,
                    departamento=departamento, provincia=provincia, distrito=distrito,
                    tipo_via=tipo_via, nombre_via=nombre_via,
                    latitud=latitud, longitud=longitud,
                )
                db.session.add(paciente)
                db.session.flush()  # asigna id sin cerrar la transacción

            estado = request.form.get("estado", "Finalizado")

            atencion = Atencion(
                paciente_id=paciente.id,
                fecha=date.today(),
                hora=datetime.now().strftime("%H:%M"),
                motivo=request.form.get("motivo", "").strip(),
                sintomas=request.form.get("sintomas", "").strip(),
                diagnostico=request.form.get("diagnostico", "").strip(),
                receta=request.form.get("receta", "").strip(),
                notas=request.form.get("notas", "").strip(),
                estado=estado,
            )
            db.session.add(atencion)
            db.session.commit()

            flash(f"Atención registrada para {paciente.nombre_completo}.", "success")
            return redirect(url_for("pacientes_hoy"))

        return render_template(
            "nueva_atencion.html",
            hoy=date.today(),
            google_maps_key=app.config["GOOGLE_MAPS_API_KEY"],
            mapa_lat_defecto=app.config["MAPA_LAT_DEFECTO"],
            mapa_lng_defecto=app.config["MAPA_LNG_DEFECTO"],
        )
    
    from datetime import date
    from flask import flash, redirect, url_for, request, render_template

    @app.route("/tareaPruebaApi", methods=["GET", "POST"])
    # ⚠️ No ponemos @login_requerido para que sea de acceso público
    def tarea_prueba_api():
        if request.method == "POST":
            # Simulamos que todo salió bien sin tocar Supabase
            flash("Registro enviado correctamente.")
            return redirect(url_for("tarea_prueba_api"))
        
        return render_template(
            "tarea_prueba.html",
            hoy=date.today(),
            nombre_clinica=app.config.get("NOMBRE_CLINICA", "Clínica (Modo Prueba)"),
            google_maps_key=app.config.get("GOOGLE_MAPS_API_KEY", ""),
            mapa_lat_defecto=app.config.get("MAPA_LAT_DEFECTO", -10.6678),
            mapa_lng_defecto=app.config.get("MAPA_LNG_DEFECTO", -76.2567)
        )

    # ---------------------------------------------------------------
    # Pacientes atendidos hoy
    # ---------------------------------------------------------------
    @app.route("/pacientes-hoy")
    @login_requerido
    def pacientes_hoy():
        hoy = date.today()
        atenciones = (
            Atencion.query.filter(Atencion.fecha == hoy)
            .order_by(Atencion.creado_en.desc())
            .all()
        )
        return render_template("pacientes_hoy.html", atenciones=atenciones, hoy=hoy)

    @app.route("/atencion/<int:atencion_id>/estado", methods=["POST"])
    @login_requerido
    def cambiar_estado(atencion_id):
        atencion = Atencion.query.get_or_404(atencion_id)
        nuevo_estado = request.form.get("estado")
        if nuevo_estado in ("Finalizado", "En proceso"):
            atencion.estado = nuevo_estado
            db.session.commit()
        return redirect(url_for("pacientes_hoy"))

    # ---------------------------------------------------------------
    # Reportes (histórico con filtro de fechas)
    # ---------------------------------------------------------------
    @app.route("/reportes")
    @login_requerido
    def reportes():
        hoy = date.today()
        desde_str = request.args.get("desde", (hoy - timedelta(days=7)).strftime("%Y-%m-%d"))
        hasta_str = request.args.get("hasta", hoy.strftime("%Y-%m-%d"))

        try:
            desde = datetime.strptime(desde_str, "%Y-%m-%d").date()
            hasta = datetime.strptime(hasta_str, "%Y-%m-%d").date()
        except ValueError:
            desde, hasta = hoy - timedelta(days=7), hoy

        atenciones = (
            Atencion.query.filter(Atencion.fecha >= desde, Atencion.fecha <= hasta)
            .order_by(Atencion.fecha.desc(), Atencion.creado_en.desc())
            .all()
        )

        total = len(atenciones)
        pacientes_unicos = len({a.paciente_id for a in atenciones})

        # Puntos para el mapa de calor: solo pacientes con ubicación guardada.
        puntos_mapa = [
            {"lat": a.paciente.latitud, "lng": a.paciente.longitud}
            for a in atenciones
            if a.paciente and a.paciente.tiene_ubicacion
        ]

        return render_template(
            "reportes.html",
            atenciones=atenciones,
            desde=desde_str,
            hasta=hasta_str,
            total=total,
            pacientes_unicos=pacientes_unicos,
            puntos_mapa=puntos_mapa,
            google_maps_key=app.config["GOOGLE_MAPS_API_KEY"],
            mapa_lat_defecto=app.config["MAPA_LAT_DEFECTO"],
            mapa_lng_defecto=app.config["MAPA_LNG_DEFECTO"],
        )

    # ---------------------------------------------------------------
    # Exportar a Excel (día actual o rango del filtro de reportes)
    # ---------------------------------------------------------------
    def _generar_excel(atenciones, titulo):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Atenciones"

        encabezados = [
            "Fecha", "Hora", "DNI", "Paciente", "Teléfono",
            "Departamento", "Provincia", "Distrito", "Dirección",
            "Latitud", "Longitud",
            "Motivo", "Síntomas", "Diagnóstico", "Receta médica", "Notas", "Estado",
        ]
        ws.append(encabezados)

        header_fill = PatternFill(start_color="1B4B43", end_color="1B4B43", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        for col in range(1, len(encabezados) + 1):
            celda = ws.cell(row=1, column=col)
            celda.fill = header_fill
            celda.font = header_font
            celda.alignment = Alignment(horizontal="center", vertical="center")

        for a in atenciones:
            p = a.paciente
            via = f"{p.tipo_via or ''} {p.nombre_via or ''}".strip() if p else ""
            ws.append([
                a.fecha.strftime("%d/%m/%Y") if a.fecha else "",
                a.hora or "",
                p.dni if p else "",
                p.nombre_completo if p else "",
                p.telefono if p else "",
                p.departamento if p else "",
                p.provincia if p else "",
                p.distrito if p else "",
                via,
                p.latitud if p and p.latitud is not None else "",
                p.longitud if p and p.longitud is not None else "",
                a.motivo or "",
                a.sintomas or "",
                a.diagnostico or "",
                a.receta or "",
                a.notas or "",
                a.estado or "",
            ])

        anchos = [12, 8, 12, 26, 14, 14, 14, 16, 22, 12, 12, 22, 22, 22, 30, 25, 14]
        for i, ancho in enumerate(anchos, start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = ancho

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    @app.route("/exportar/hoy")
    @login_requerido
    def exportar_hoy():
        hoy = date.today()
        atenciones = Atencion.query.filter(Atencion.fecha == hoy).order_by(Atencion.hora).all()
        buffer = _generar_excel(atenciones, f"Atenciones {hoy.strftime('%d-%m-%Y')}")
        return send_file(
            buffer,
            as_attachment=True,
            download_name=f"atenciones_{hoy.strftime('%Y%m%d')}.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    @app.route("/exportar/rango")
    @login_requerido
    def exportar_rango():
        desde_str = request.args.get("desde")
        hasta_str = request.args.get("hasta")
        desde = datetime.strptime(desde_str, "%Y-%m-%d").date()
        hasta = datetime.strptime(hasta_str, "%Y-%m-%d").date()

        atenciones = (
            Atencion.query.filter(Atencion.fecha >= desde, Atencion.fecha <= hasta)
            .order_by(Atencion.fecha, Atencion.hora)
            .all()
        )
        buffer = _generar_excel(atenciones, f"Atenciones {desde_str} a {hasta_str}")
        return send_file(
            buffer,
            as_attachment=True,
            download_name=f"reporte_{desde_str}_a_{hasta_str}.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    @app.context_processor
    def inject_globals():
        return {"nombre_clinica": app.config["NOMBRE_CLINICA"]}

    return app




app = create_app()

if __name__ == "__main__":
    app.run(debug=True, host="0.0.0.0", port=5000)
